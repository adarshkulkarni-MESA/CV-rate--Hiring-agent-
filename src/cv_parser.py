"""Reads a CV file (.docx, .pdf, .txt) into plain text.

Kept deliberately dumb: this is extraction, not understanding. Understanding
happens in scorer.py, against redacted text only.
"""
from pathlib import Path


def parse_cv(path: str) -> str:
    p = Path(path)
    suffix = p.suffix.lower()

    if suffix == ".docx":
        return _parse_docx(p)
    if suffix == ".pdf":
        return _parse_pdf(p)
    if suffix in (".txt", ".md"):
        return p.read_text(encoding="utf-8", errors="ignore")

    raise ValueError(
        f"Unsupported CV format: {suffix}. "
        "This build handles .docx, .pdf, and .txt — the case's applications/ "
        "folder is described as 'mixed formats,' so extend here if it includes others."
    )


def _parse_docx(p: Path) -> str:
    import docx

    doc = docx.Document(str(p))
    lines = [para.text for para in doc.paragraphs if para.text.strip()]

    # tables (some CVs put skills/experience in table cells)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines)


def _parse_pdf(p: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(p))
    return "\n".join(page.extract_text() or "" for page in reader.pages)
